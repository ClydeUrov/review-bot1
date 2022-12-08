import requests
from bs4 import BeautifulSoup as BS
from openpyxl import load_workbook
import telegram
import time


def fill_table(wb, html_selector, name, page):
	sheet = wb[page]
	sheet.delete_cols(0, 100)
	sheet.append(['', name])
	sheet.append(['№', 'Название', 'Описание', 'Ссылка'])
	number = 0
	if page == 'List1':
		description = 'article'
		http = ''
	elif page == 'List2':
		description = 'span'
		http = 'https://www.weblancer.net'
	for element in html_selector:
		number += 1
		sheet.append([
			number,
			element.find('a').text, 
			element.find(description).text, 
			http + element.find('a').get('href')
		])


if __name__ == '__main__':
	bot = telegram.Bot(token="5929423015:AAFU8to5_VT2-ffHxplJGRc0KT0Y3MDNKCk")
	wb = load_workbook('example.xlsx')
	name = 'python'
	params = {
		'q': name
	}
	view_html = requests.get('https://freelance.ua/orders/', params=params)
	html = BS(view_html.content, 'html.parser')
	fill_table(wb, html.select(".container > .media-body"), 'Freelance', 'List1')


	view_html = requests.get('https://www.weblancer.net/jobs/python/?lang=ua')
	html = BS(view_html.content, 'html.parser')
	fill_table(wb, html.select(".click_container-link")[1:], 'Weblancer', 'List2')


	wb.save('example.xlsx')
	wb.close()
	bot.sendDocument(
		chat_id="513965977",
		document=open('./example.xlsx', 'rb')
	)
	time.sleep(604800)
